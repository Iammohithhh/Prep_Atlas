"""Private, resumable OCR index. Nothing here is copied into the public site.

OCR is a reading aid; it is never automatically published as a question.
Run: python build/extract_sources.py [--company NAME] [--limit N]
"""
import argparse
import hashlib
import json
import os
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / 'build' / 'ocr'


def read_cache(path, rel):
    try:
        record = json.loads(path.read_text(encoding='utf-8'))
        if isinstance(record, dict) and record.get('file') == rel and isinstance(record.get('text'), str):
            return record
    except (OSError, UnicodeError, json.JSONDecodeError):
        pass
    return None


def write_json_atomic(path, value):
    temporary = path.with_name(path.name + f'.{os.getpid()}.tmp')
    try:
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--company')
    parser.add_argument('--limit', type=int)
    parser.add_argument('--audit', action='store_true', help='Check cache coverage without running OCR.')
    parser.add_argument('--retry-errors', action='store_true', help='Re-extract cached records that contain errors.')
    args = parser.parse_args()
    manifest = json.loads((ROOT / 'build/manifest.json').read_text(encoding='utf-8'))
    CACHE.mkdir(exist_ok=True)
    selected = {c: paths for c, paths in manifest.items() if not args.company or c.lower() == args.company.lower()}
    if not selected:
        parser.error('Company not found in manifest.')
    if args.limit is not None and args.limit < 1:
        parser.error('--limit must be positive.')
    total = sum(map(len, selected.values()))
    if args.audit:
        valid = errors = missing = invalid = 0
        for files in selected.values():
            for rel in files:
                dest = CACHE / (hashlib.sha256(rel.encode()).hexdigest()[:20]+'.json')
                record = read_cache(dest, rel)
                if record is not None:
                    valid += 1
                    errors += bool(record.get('error'))
                elif dest.exists(): invalid += 1
                else: missing += 1
        print(f'Cache audit: {valid}/{total} valid records; {errors} extraction errors; {missing} missing; {invalid} invalid.', flush=True)
        return
    ocr = None
    index_path = CACHE / ('index-'+hashlib.sha256(args.company.encode()).hexdigest()[:8]+'.json' if args.company else 'index.json')
    index = {}
    done = 0
    for company, files in selected.items():
        for rel in files:
            key = hashlib.sha256(rel.encode()).hexdigest()[:20]
            dest = CACHE / (key + '.json')
            cached = read_cache(dest, rel)
            if cached is not None and not (args.retry_errors and cached.get('error')):
                index[rel] = dest.name
                continue
            claim = dest.with_suffix('.lock')
            try:
                claim_fd = os.open(claim, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                with os.fdopen(claim_fd, 'w') as claim_file:
                    claim_file.write(str(os.getpid()))
            except FileExistsError:
                continue
            path = ROOT / rel
            record = {'company': company, 'file': rel, 'text': '', 'confidence': None}
            completed = False
            try:
                if ocr is None:
                    from rapidocr_onnxruntime import RapidOCR
                    ocr = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)
                suffix = path.suffix.lower()
                if suffix == '.pdf':
                    import pymupdf
                    with pymupdf.open(path) as pdf:
                        parts = []
                        for page in pdf:
                            txt = page.get_text()
                            if len(txt.strip()) < 60:
                                pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5))
                                import numpy as np
                                arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                                rows, _ = ocr(arr[:, :, :3])
                                txt = '\n'.join(row[1] for row in rows or [])
                            parts.append(txt)
                        record['text'] = '\n\n--- PAGE ---\n\n'.join(parts)
                elif suffix == '.docx':
                    from docx import Document
                    doc = Document(path)
                    record['text'] = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
                elif suffix == '.txt':
                    record['text'] = path.read_text(encoding='utf-8', errors='replace')
                else:
                    import numpy as np
                    from PIL import Image, ImageOps
                    with Image.open(path) as im:
                        im = ImageOps.exif_transpose(im).convert('RGB')
                        im.thumbnail((1800, 2400))
                        rows, _ = ocr(np.asarray(im))
                    record['text'] = '\n'.join(row[1] for row in rows or [])
                    record['confidence'] = round(sum(row[2] for row in rows) / len(rows), 3) if rows else 0
                completed = True
            except Exception as exc:
                record['error'] = str(exc)
                completed = True
            finally:
                try:
                    if completed:
                        write_json_atomic(dest, record)
                finally:
                    claim.unlink(missing_ok=True)
            index[rel] = dest.name
            write_json_atomic(index_path, index)
            done += 1
            if done % 10 == 0:
                print(f'{len(index)} / {total} indexed; {company}', flush=True)
            if args.limit and done >= args.limit:
                return
    write_json_atomic(index_path, index)
    print(f'Extraction complete: {len(index)} / {total} cached files; {done} extracted this run.', flush=True)


if __name__ == '__main__':
    main()
